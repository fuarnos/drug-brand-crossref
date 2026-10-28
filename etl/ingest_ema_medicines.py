# -*- coding: utf-8 -*-
"""EMA 集中審批藥品數據（EU 層）入庫準備。

數據：`data/raw/pt/ema_medicines_report_en.xlsx`
  - 官方下載頁：https://www.ema.europa.eu/en/medicines/download-medicine-data
  - 本次快照：2,754 行 = 2,754 個 EMA product number（1 行 1 個 EU 號）；
    Human 2,357 / Veterinary 397；狀態 Authorised 1,875（其餘 withdrawn/refused/…）。
  - **EMA XLSX 沒有 strength（規格）與 pharmaceutical form（劑型）欄位**，名稱只是
    品牌名（如 "Cirbloc"），不能可靠解析規格/劑型；EU 層只能做「品牌名 → 活性成分 /
    MAH / ATC / 狀態」對照，不能做 D1 規格對應。

許可（EMA Legal notice，2026-10-28 快照原文）：
  "Information and documents made available on EMA's webpages are public and may be
   reproduced and/or distributed, totally or in part, irrespective of the means and/or
   the formats used, for non-commercial and commercial purposes, provided that EMA is
   always acknowledged as the source of the material. Such acknowledgement must be
   included in each copy of the material."
  → 可再發布（含商業），必須署名 EMA；source.license 建議寫
    "© European Medicines Agency; re-use permitted with acknowledgement (EMA legal notice)"。

EU 層口徑（**不是國家市場**）：EMA 集中審批授權在 EU/EEA 全境有效，屬 region=EU；
正式寫入用 **product.region='EU' + product.country=NULL**（db/schema.sql 已加；XOR CHECK），
不進國家下拉/國家統計；`coverage` 視圖用 COALESCE(country, region) 給 EU 單獨一行。
兩種寫入模式：
  - `--apply`（staging 原型表）保留，供兼容/比對；
  - **`--apply-products`（正式 product/product_component/name 表，region='EU'）**：
    需要先跑 `db/init_db.py` 完成 region schema 遷移；單事務 + 撤銷清單 + 幂等
    （source_id='ema_medicines' 先刪後寫）；`--rollback-products --manifest ...` 撤銷。
  EMA 數據無 strength/form/pack，product 寫 `dosage_form='unknown'`、`release='none'`、
  strength NULL；頁面須明示「EU 集中審批（EMA），非國別產品，無規格/劑型/包裝」。

CLI：
    python etl/ingest_ema_medicines.py --dry-run                 # 解析+映射+統計（默認）
    python etl/ingest_ema_medicines.py --apply --db .tmp/ema-ingest/test.db      # staging
    python etl/ingest_ema_medicines.py --rollback --db .tmp/ema-ingest/test.db
    python etl/ingest_ema_medicines.py --apply-products --db <副本> --allow-real # 正式表
    python etl/ingest_ema_medicines.py --rollback-products --manifest <json> --db <副本>
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "etl"))
from ingest_openfda_ndc import slug, split_salt  # noqa: E402
from normalize_name import normalize  # noqa: E402
from normalize_salts import classify_salt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

DEFAULT_XLSX = ROOT / "data" / "raw" / "pt" / "ema_medicines_report_en.xlsx"
DEFAULT_DB = ROOT / "data" / "drugs.db"
WORK = ROOT / ".tmp" / "ema-ingest"
SOURCE_ID = "ema_medicines"
REGION = "EU"
SOURCE_URL = "https://www.ema.europa.eu/en/medicines/download-medicine-data"
SOURCE_LICENSE = "© European Medicines Agency; re-use permitted with acknowledgement (EMA legal notice)"
SOURCE_LICENSE_URL = "https://www.ema.europa.eu/en/about-us/about-website/legal-notice"
DATA_AS_OF = "2026-10-28"
TODAY = date.today().isoformat()
PRODUCT_MANIFEST = WORK / "ema_products_manifest.json"

EXTRA_SALT_SUFFIXES = (
    "besilate", "mesilate", "besylate", "mesylate", "tosilate", "tosylate",
    "etexilate", "hyclate", "estolate", "glucoheptonate", "lactobionate",
    "dihydrochloride", "hydrobromide", "hydrochloride", "sulfate", "sulphate",
    "phosphate", "acetate", "fumarate", "maleate", "succinate", "tartrate",
    "citrate", "lactate", "malate", "aspartate", "gluconate", "benzoate",
    "pamoate", "embonate", "xinafoate", "furoate", "propionate", "valerate",
    "palmitate", "decanoate", "enanthate", "cypionate", "undecanoate",
    "sodium", "potassium", "calcium", "magnesium", "chloride", "bromide",
    "iodide", "oxide", "hydroxide", "carbonate", "bicarbonate",
)
HYDRATE_WORDS = {
    "monohydrate", "dihydrate", "trihydrate", "tetrahydrate", "hemihydrate",
    "sesquihydrate", "hydrate", "anhydrous", "solvate", "ethanolate",
}

STATUS_MAP = {
    "Authorised": "active",
    "Withdrawn": "withdrawn",
    "Revoked": "withdrawn",
    "Lapsed": "withdrawn",
    "Expired": "withdrawn",
    "Suspended": "discontinued",
    "Refused": "unknown",
    "Application withdrawn": "unknown",
    "Opinion": "unknown",
    "Withdrawn from rolling review": "unknown",
}

STAGING_DDL = """
CREATE TABLE IF NOT EXISTS ema_medicine (
    eu_number          TEXT PRIMARY KEY,
    medicine_name      TEXT NOT NULL,
    category           TEXT,
    status_raw         TEXT,
    status             TEXT,
    atc                TEXT,
    mah                TEXT,
    therapeutic_area   TEXT,
    medicine_url       TEXT,
    authorisation_date TEXT,
    first_published    TEXT,
    last_updated       TEXT,
    region             TEXT NOT NULL DEFAULT 'EU',
    source_id          TEXT NOT NULL DEFAULT 'ema_medicines',
    n_ingredients      INTEGER NOT NULL DEFAULT 0,
    ingested_at        TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS ema_medicine_ingredient (
    eu_number         TEXT NOT NULL,
    ord               INTEGER NOT NULL,
    active_substance  TEXT NOT NULL,
    substance_id      TEXT,
    mapping_method    TEXT NOT NULL,
    PRIMARY KEY (eu_number, ord)
);
CREATE INDEX IF NOT EXISTS idx_ema_ing_substance
    ON ema_medicine_ingredient (substance_id);
CREATE INDEX IF NOT EXISTS idx_ema_medicine_status
    ON ema_medicine (status, category);
"""


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _split_ingredients(text: str) -> list[str]:
    """活性成分串 → token；只按 / ; + 和 " and " 切（逗号是描述的一部分，不能切）。"""
    text = (text or "").strip()
    if not text or text.lower() in ("none", "-"):
        return []
    # 只按「帶空格的 /」「+」「;」「 and 」切；不切疫苗株名的緊湊斜線
    # （A/California/7/2009 不能被拆成 California / 7）；純數字 token 丟棄。
    parts = re.split(r"\s+/\s+|\s*[+;]\s*|\s+and\s+", text)
    out = []
    for p in parts:
        t = p.strip(" .,;")
        if not t or t.lower() in ("none", "-"):
            continue
        if re.fullmatch(r"[\d.,\s]+", t):
            continue
        out.append(t)
    return out


def _strip_salt_words(token: str) -> list[str]:
    """由含鹽/水合物的成分名推導可匹配 base 名（含英式拼法）。"""
    words = token.lower().replace("-", " ").split()
    candidates = [" ".join(words)]
    changed = True
    while changed and len(words) > 1:
        changed = False
        for suf in EXTRA_SALT_SUFFIXES + tuple(HYDRATE_WORDS):
            if words and words[-1] == suf:
                words = words[:-1]
                candidates.append(" ".join(words))
                changed = True
                break
    if len(words) > 1:
        candidates.append(" ".join(words[:-1]))
    return list(dict.fromkeys(c for c in candidates if c))


class SubstanceMapper:
    """EMA 成分名 → 現有 substance.substance_id（slug / INN 正規化 / 拆鹽 base）。"""

    def __init__(self, db: sqlite3.Connection):
        self.by_id = {r[0] for r in db.execute("SELECT substance_id FROM substance")}
        self.by_inn = {}
        for sid, inn in db.execute("SELECT substance_id, inn_en FROM substance"):
            self.by_inn.setdefault(_norm(inn), sid)

    def map(self, token: str) -> tuple[str, str]:
        token = token.strip()
        if not token:
            return "", "empty"
        sid = slug(token)
        if sid in self.by_id:
            return sid, "slug"
        n = _norm(token)
        if n in self.by_inn:
            return self.by_inn[n], "inn"
        base, salt = split_salt(token)
        if salt:
            b = slug(base)
            if b in self.by_id:
                return b, "salt-strip"
        for cand in _strip_salt_words(token):
            b = slug(cand)
            if b in self.by_id:
                return b, "salt-normalize"
            bn = _norm(cand)
            if bn in self.by_inn:
                return self.by_inn[bn], "salt-normalize-inn"
        clean = re.sub(r"\([^)]*\)", " ", token)
        clean = re.sub(r"\d+(?:[.,]\d+)?\s*(?:mg|g|mcg|µg|iu|ml|%)", " ", clean, flags=re.I)
        b = slug(clean)
        if b in self.by_id:
            return b, "cleanup"
        return "", "unmapped"


def parse_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = next(i for i, r in enumerate(rows)
                 if any(str(c).strip() == "Name of medicine" for c in r if c))
    header = [str(c).strip() if c is not None else "" for c in rows[hdr_i]]
    while header and header[-1] == "":
        header.pop()
    out = []
    for r in rows[hdr_i + 1:]:
        vals = list(r[:len(header)])
        if not any(v not in (None, "") for v in vals):
            continue
        out.append({header[i]: vals[i] for i in range(len(header))})
    return out


def _val(row: dict, key: str, default: str = "") -> str:
    v = row.get(key)
    if v in (None, ""):
        return default
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def build_records(rows: list[dict], mapper: SubstanceMapper,
                  category: str = "human") -> tuple[list[dict], list[dict], dict]:
    meds, comps = [], []
    stats: dict[str, int] = {}

    def bump(key: str, n: int = 1) -> None:
        stats[key] = stats.get(key, 0) + n

    for row in rows:
        cat = _val(row, "Category").lower() or "unknown"
        if category != "all" and cat != category:
            bump("skip_category")
            continue
        eu = _val(row, "EMA product number")
        name = _val(row, "Name of medicine")
        if not eu or not name:
            bump("skip_no_key")
            continue
        status_raw = _val(row, "Medicine status")
        status = STATUS_MAP.get(status_raw, "unknown")
        ingredients = _split_ingredients(_val(row, "Active substance"))
        if not ingredients:
            ingredients = _split_ingredients(
                _val(row, "International non-proprietary name (INN) / common name"))
        mapped = [(t, *mapper.map(t)) for t in ingredients]
        n_map = sum(1 for _, sid, _m in mapped if sid)
        bump("rows")
        bump(f"status:{status_raw or '(empty)'}")
        bump(f"mapped_status:{status}")
        bump("ingredients", len(mapped))
        bump("ingredients_mapped", n_map)
        if not mapped:
            bump("rows_no_ingredient")
        elif n_map == 0:
            bump("rows_no_mapped_ingredient")
        if mapped and n_map == len(mapped):
            bump("rows_all_mapped")
        meds.append({
            "eu_number": eu, "medicine_name": name, "category": cat,
            "status_raw": status_raw, "status": status,
            "atc": _val(row, "ATC code (human)") or _val(row, "ATCvet code (veterinary)"),
            "mah": _val(row, "Marketing authorisation developer / applicant / holder"),
            "therapeutic_area": _val(row, "Therapeutic area (MeSH)"),
            "medicine_url": _val(row, "Medicine URL"),
            "authorisation_date": _val(row, "Marketing authorisation date"),
            "first_published": _val(row, "First published date"),
            "last_updated": _val(row, "Last updated date"),
            "n_ingredients": len(mapped),
            "mapping_methods": sorted({m for _, _s, m in mapped}),
        })
        for ord_, (token, sid, method) in enumerate(mapped):
            bump(f"map_method:{method}")
            comps.append({"eu_number": eu, "ord": ord_, "active_substance": token,
                          "substance_id": sid, "mapping_method": method})
    return meds, comps, stats


def summarize(meds: list[dict], comps: list[dict], stats: dict) -> dict:
    unmapped = Counter(c["active_substance"] for c in comps if not c["substance_id"])
    mapped_tokens = {c["active_substance"] for c in comps if c["substance_id"]}
    all_tokens = {c["active_substance"] for c in comps}
    mapped_substances = {c["substance_id"] for c in comps if c["substance_id"]}
    mapped_by_eu = {c["eu_number"] for c in comps if c["substance_id"]}
    human = [m for m in meds if m["category"] == "human"]
    authorised = [m for m in human if m["status"] == "active"]
    n_cmp = max(1, len(comps))
    n_mapped_cmp = sum(1 for c in comps if c["substance_id"])
    return {
        "rows_used": len(meds),
        "human_rows": len(human),
        "human_authorised": len(authorised),
        "unique_eu": len({m["eu_number"] for m in meds}),
        "components": len(comps),
        "unique_ingredient_tokens": len(all_tokens),
        "mapped_ingredient_tokens": len(mapped_tokens),
        "ingredient_token_rate": round(len(mapped_tokens) / max(1, len(all_tokens)) * 100, 1),
        "mapped_components": n_mapped_cmp,
        "component_rate": round(n_mapped_cmp / n_cmp * 100, 1),
        "rows_with_mapped_ingredient": sum(1 for m in meds if m["eu_number"] in mapped_by_eu),
        "rows_no_ingredient": stats.get("rows_no_ingredient", 0),
        "rows_no_mapped": stats.get("rows_no_mapped_ingredient", 0),
        "rows_all_mapped": stats.get("rows_all_mapped", 0),
        "map_methods": {k.split(":", 1)[1]: v for k, v in stats.items()
                        if k.startswith("map_method:")},
        "statuses": {k.split(":", 1)[1]: v for k, v in stats.items()
                     if k.startswith("status:")},
        "distinct_substances_mapped": len(mapped_substances),
        "unmapped_top": unmapped.most_common(40),
        "unmapped_unique": len(unmapped),
    }


def _connect(db_path: Path, write: bool) -> sqlite3.Connection:
    if write:
        return sqlite3.connect(db_path, timeout=120)
    return sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True, timeout=120)


def _guard_real(db_path: Path, action: str, allow_real: bool) -> None:
    try:
        is_real = db_path.resolve() == DEFAULT_DB.resolve()
    except OSError:
        is_real = False
    if is_real and action == "apply" and not allow_real:
        raise SystemExit(
            f"拒絕在真實庫 {db_path} 寫入 EMA 原型表：本任務只允許 DB 副本。\n"
            "確需寫真實庫時加 --allow-real（仍建議先備份；正式接入應先完成 region schema 變更）。")


def _write_preview(meds, comps, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    by_med: dict[str, list[dict]] = {}
    for c in comps:
        by_med.setdefault(c["eu_number"], []).append(c)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["eu_number", "medicine_name", "category", "status_raw", "status", "atc",
                    "mah", "active_substances", "substance_ids", "mapping_methods",
                    "therapeutic_area", "authorisation_date", "medicine_url"])
        for m in meds:
            cs = by_med.get(m["eu_number"], [])
            w.writerow([m["eu_number"], m["medicine_name"], m["category"],
                        m["status_raw"], m["status"], m["atc"], m["mah"],
                        " | ".join(c["active_substance"] for c in cs),
                        " | ".join(c["substance_id"] or "" for c in cs),
                        " | ".join(c["mapping_method"] for c in cs),
                        m["therapeutic_area"], m["authorisation_date"], m["medicine_url"]])


def _require_region_schema(db: sqlite3.Connection) -> None:
    row = db.execute("SELECT sql FROM sqlite_master WHERE name='product'").fetchone()
    if not row or "region" not in (row[0] or ""):
        raise SystemExit(
            "product.region 不存在：先跑 `python db/init_db.py --db <db>` 完成 region schema 遷移"
            "（country 可空 + region 列 + XOR CHECK + coverage COALESCE）。")


def _ensure_region_source(db: sqlite3.Connection) -> None:
    if db.execute("SELECT 1 FROM source WHERE source_id=?", (SOURCE_ID,)).fetchone():
        return
    db.execute(
        "INSERT INTO source (source_id, name, url, license, license_url, retrieved_at,"
        " data_as_of, data_as_of_note, notes) VALUES (?,?,?,?,?,?,?,?,?)",
        (SOURCE_ID, "European Medicines Agency — Medicines data", SOURCE_URL,
         SOURCE_LICENSE, SOURCE_LICENSE_URL, TODAY, DATA_AS_OF,
         "XLSX 按需生成（EMA download page）；本快照 2026-10-28",
         "EU 集中審批（centralised procedure）數據；再發布須每份內容署名 EMA。"))


def _variant_resolver(db: sqlite3.Connection):
    """(substance_id, salt token) → variant_id 解析器（含 salt-orphan 兜底）。"""
    vidx: dict[tuple[str, str], str] = {}
    by_sid: dict[str, list[tuple[str, str]]] = {}
    for r in db.execute("SELECT variant_id, substance_id, COALESCE(salt_code,'') sc FROM variant"):
        vidx[(r["substance_id"], r["sc"])] = r["variant_id"]
        by_sid.setdefault(r["substance_id"], []).append((r["variant_id"], r["sc"]))
    salt_codes = {sc for (_sid, sc) in vidx if sc}

    def resolve(sid: str, token: str) -> str | None:
        _base, salt = split_salt(token)
        sc = ""
        if salt:
            try:
                cl = classify_salt(salt)
                sc = cl.salt_code if cl.kind == "salt" else ""
            except Exception:  # noqa: BLE001
                sc = ""
        vid = vidx.get((sid, sc)) or vidx.get((sid, ""))
        if vid:
            return vid
        cands = by_sid.get(sid) or []
        if len(cands) == 1:
            return cands[0][0]
        for v, s in cands:
            if s and (s == sc or s in sid):
                return v
        if cands:
            return next((v for v, s in cands if s), cands[0][0])
        for one in sorted(salt_codes, key=len, reverse=True):
            suffix = "-" + one
            if sid.endswith(suffix):
                base = sid[:-len(suffix)]
                bvid = vidx.get((base, one)) or vidx.get((base, ""))
                if bvid:
                    return bvid
                for v, _s in (by_sid.get(base) or []):
                    return v
        return None

    return resolve


def cmd_apply_products(args: argparse.Namespace) -> int:
    """EMA human → 正式 product/product_component/name（region='EU'，country=NULL）。"""
    import time
    _guard_real(args.db, "apply", args.allow_real)
    if not args.db.exists():
        raise SystemExit(f"DB 不存在：{args.db}")
    rows = parse_xlsx(args.xlsx)
    db = sqlite3.connect(args.db, timeout=120)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    manifest_path = args.manifest
    if not manifest_path.is_absolute():
        manifest_path = ROOT / manifest_path
    prev_keys: list[str] = []
    prev_variants: list[list[str]] = []
    if manifest_path.exists():
        try:
            old = json.loads(manifest_path.read_text(encoding="utf-8"))
            if old.get("source_id") == SOURCE_ID:
                prev_keys = list(old.get("created_name_keys") or [])
                prev_variants = [list(x) for x in (old.get("created_name_variants") or [])]
        except (OSError, ValueError):
            pass
    t0 = time.perf_counter()
    try:
        _require_region_schema(db)
        _ensure_region_source(db)
        db.commit()
        mapper = SubstanceMapper(db)
        meds, comps, stats = build_records(rows, mapper, args.category)
        summary = summarize(meds, comps, stats)
        resolve = _variant_resolver(db)
        comps_by_eu: dict[str, list[dict]] = {}
        for c in comps:
            comps_by_eu.setdefault(c["eu_number"], []).append(c)
        status_counts: dict[str, int] = {}
        status_map = {"active": "active", "discontinued": "discontinued",
                      "withdrawn": "withdrawn", "unknown": "unknown"}
        db.execute("BEGIN IMMEDIATE")
        removed = db.execute("DELETE FROM product WHERE region=? AND source_id=?",
                             (REGION, SOURCE_ID)).rowcount
        created_keys: list[str] = []
        created_variants: list[list[str]] = []
        touched: set[str] = set()
        product_ids: list[str] = []
        n_components = n_no_variant = n_no_component = 0
        run_id = f"{SOURCE_ID}-{TODAY}"
        for med in meds:
            eu = med["eu_number"]
            pid = "EU-" + (slug(eu) or "x")[:48]
            if db.execute("SELECT source_id FROM product WHERE product_id=?", (pid,)).fetchone():
                pid = f"{pid}-{abs(hash(eu)) % 100000:05d}"
            status = status_map.get(med["status"], "unknown")
            status_counts[status] = status_counts.get(status, 0) + 1
            note = " · ".join(x for x in [
                "EU region layer（EMA centralised procedure）— 非國別產品",
                f"EMA status: {med['status_raw'] or '?'}",
                f"ATC: {med['atc']}" if med.get("atc") else "",
                f"Therapeutic area: {med['therapeutic_area']}" if med.get("therapeutic_area") else "",
                f"Authorisation date: {med['authorisation_date']}" if med.get("authorisation_date") else "",
                "無規格/劑型/包裝（EMA 數據不含）",
            ] if x)[:500]
            db.execute(
                "INSERT OR REPLACE INTO product (product_id, country, region, dosage_form,"
                " form_base, route, release, pack, mah, approval_id, category, is_branded,"
                " status, status_date, source_id, source_url, retrieved_at, note)"
                " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (pid, None, REGION, "unknown", None, None, "none", None,
                 med.get("mah") or None, eu or None, "EMA centralised (human)", 1,
                 status, med.get("authorisation_date") or None, SOURCE_ID,
                 med.get("medicine_url") or SOURCE_URL, TODAY, note))
            product_ids.append(pid)
            raw_name = (med.get("medicine_name") or "").strip()
            if raw_name:
                nm = normalize(raw_name)
                key = nm.norm
                if key:
                    script = nm.script if nm.script in (
                        "latin", "cyrillic", "han", "kana", "hangul") else "latin"
                    cur = db.execute(
                        "INSERT OR IGNORE INTO name_key (key_id, norm, compact, script)"
                        " VALUES (?,?,?,?)", (key, key, key.replace(" ", ""), script))
                    if cur.rowcount:
                        created_keys.append(key)
                    touched.add(key)
                    cur2 = db.execute(
                        "INSERT OR IGNORE INTO name_variant (key_id, raw, source_id)"
                        " VALUES (?,?,?)", (key, raw_name, SOURCE_ID))
                    if cur2.rowcount:
                        created_variants.append([key, raw_name])
                    db.execute(
                        "INSERT OR IGNORE INTO product_name (product_id, key_id, name_type,"
                        " is_primary) VALUES (?,?,?,?)", (pid, key, "trade", 1))
            ord_ = 0
            for c in comps_by_eu.get(eu, []):
                sid = c["substance_id"]
                if not sid:
                    continue
                vid = resolve(sid, c["active_substance"])
                if not vid:
                    n_no_variant += 1
                    continue
                db.execute(
                    "INSERT OR REPLACE INTO product_component (product_id, ord, variant_id,"
                    " strength_value, strength_unit, per_volume_value, per_volume_unit)"
                    " VALUES (?,?,?,?,?,?,?)", (pid, ord_, vid, None, None, None, None))
                ord_ += 1
                n_components += 1
            if ord_ == 0:
                n_no_component += 1
        db.execute(
            "INSERT OR REPLACE INTO ingest_run (run_id, source_id, started_at, finished_at,"
            " rows_in, rows_kept, rows_rejected, status, note) VALUES (?,?,?,?,?,?,?,?,?)",
            (run_id, SOURCE_ID, TODAY, datetime.now(timezone.utc).isoformat(timespec="seconds"),
             len(meds), len(product_ids), len(meds) - len(product_ids), "ok",
             json.dumps({"region": REGION, "layer": "product-region", "summary": summary},
                        ensure_ascii=False)))
        db.commit()
        elapsed = round(time.perf_counter() - t0, 2)
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    all_keys = sorted(set(prev_keys) | set(created_keys))
    all_variants = sorted({tuple(x) for x in (prev_variants + created_variants)})
    res = {"db": str(args.db), "run_id": run_id, "human_medicines": len(meds),
           "products": len(product_ids), "components_written": n_components,
           "components_no_variant": n_no_variant, "products_without_component": n_no_component,
           "created_name_keys": len(created_keys), "removed_old_products": removed,
           "status_counts": status_counts, "elapsed_s": elapsed, "summary": summary}
    manifest = {"tool": "etl/ingest_ema_medicines.py", "db": str(args.db),
                "source_id": SOURCE_ID, "region": REGION, "created_at": TODAY,
                "product_ids": product_ids, "created_name_keys": all_keys,
                "created_name_variants": [list(x) for x in all_variants], "run_id": run_id}
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    WORK.mkdir(parents=True, exist_ok=True)
    (WORK / "products_apply_stats.json").write_text(json.dumps(res, ensure_ascii=False, indent=1),
                                                    encoding="utf-8")
    print("=== EMA EU 層（正式 product 表）寫入完成 ===")
    print(json.dumps(res, ensure_ascii=False, indent=1))
    print("撤銷清單：", manifest_path)
    return 0


def cmd_rollback_products(args: argparse.Namespace) -> int:
    _guard_real(args.db, "apply", args.allow_real)
    manifest_path = args.manifest
    if not manifest_path.is_absolute():
        manifest_path = ROOT / manifest_path
    if not manifest_path.exists():
        raise SystemExit(f"找不到撤銷清單：{manifest_path}")
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    db = sqlite3.connect(args.db, timeout=120)
    db.execute("PRAGMA foreign_keys=ON")
    try:
        db.execute("BEGIN IMMEDIATE")
        n = db.execute("DELETE FROM product WHERE source_id=? AND region=?",
                       (SOURCE_ID, REGION)).rowcount
        db.execute("DELETE FROM ingest_run WHERE run_id=?", (data.get("run_id") or "",))
        removed_keys = removed_variants = 0
        for key, raw in (data.get("created_name_variants") or []):
            cur = db.execute("DELETE FROM name_variant WHERE key_id=? AND raw=?", (key, raw))
            removed_variants += cur.rowcount
        for key in data.get("created_name_keys") or []:
            if not db.execute("SELECT 1 FROM product_name WHERE key_id=? LIMIT 1",
                              (key,)).fetchone():
                cur = db.execute("DELETE FROM name_variant WHERE key_id=?", (key,))
                removed_variants += cur.rowcount
                db.execute("DELETE FROM name_key WHERE key_id=?", (key,))
                removed_keys += 1
        db.commit()
    finally:
        db.close()
    res = {"removed_products": n, "removed_name_keys": removed_keys,
           "removed_name_variants": removed_variants, "db": str(args.db)}
    print("=== EMA EU 層回滾完成 ===", res)
    return 0


def cmd_dry_run(args: argparse.Namespace) -> int:
    rows = parse_xlsx(args.xlsx)
    db = _connect(args.db, write=False)
    try:
        mapper = SubstanceMapper(db)
        meds, comps, stats = build_records(rows, mapper, args.category)
        summary = summarize(meds, comps, stats)
    finally:
        db.close()
    summary["source"] = str(args.xlsx)
    summary["generated_at"] = datetime.now(timezone.utc).isoformat()
    WORK.mkdir(parents=True, exist_ok=True)
    (WORK / "stats.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    _write_preview(meds, comps, args.csv)
    print("=== EMA ingest dry-run ===")
    print(f"XLSX rows used : {summary['rows_used']:,}（human {summary['human_rows']:,}，"
          f"其中 Authorised {summary['human_authorised']:,}）")
    print(f"unique EU 號   : {summary['unique_eu']:,}；成分組件 {summary['components']:,}")
    print(f"成分 token     : {summary['unique_ingredient_tokens']:,}，"
          f"映射 {summary['mapped_ingredient_tokens']:,} "
          f"({summary['ingredient_token_rate']}%)")
    print(f"組件映射率     : {summary['mapped_components']:,}/{summary['components']:,} "
          f"({summary['component_rate']}%)")
    print(f"映射方法       : {summary['map_methods']}")
    print(f"狀態分布       : {summary['statuses']}")
    print(f"unmapped top   : {[x[0] for x in summary['unmapped_top'][:10]]}")
    print(f"預覽 CSV       : {args.csv}（{len(meds):,} 行）")
    print(f"統計           : {WORK / 'stats.json'}")
    return 0


def cmd_apply(args: argparse.Namespace) -> int:
    _guard_real(args.db, "apply", args.allow_real)
    if not args.db.exists():
        raise SystemExit(f"DB 不存在：{args.db}")
    rows = parse_xlsx(args.xlsx)
    db = sqlite3.connect(args.db, timeout=120)
    try:
        db.executescript(STAGING_DDL)
        mapper = SubstanceMapper(db)
        meds, comps, stats = build_records(rows, mapper, args.category)
        summary = summarize(meds, comps, stats)
        now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        with db:
            db.execute("DELETE FROM ema_medicine_ingredient")
            db.execute("DELETE FROM ema_medicine")
            db.executemany(
                """INSERT INTO ema_medicine
                   (eu_number, medicine_name, category, status_raw, status, atc, mah,
                    therapeutic_area, medicine_url, authorisation_date, first_published,
                    last_updated, region, source_id, n_ingredients, ingested_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                [(m["eu_number"], m["medicine_name"], m["category"], m["status_raw"],
                  m["status"], m["atc"], m["mah"], m["therapeutic_area"],
                  m["medicine_url"], m["authorisation_date"], m["first_published"],
                  m["last_updated"], REGION, SOURCE_ID, m["n_ingredients"], now)
                 for m in meds])
            db.executemany(
                """INSERT INTO ema_medicine_ingredient
                   (eu_number, ord, active_substance, substance_id, mapping_method)
                   VALUES (?,?,?,?,?)""",
                [(c["eu_number"], c["ord"], c["active_substance"], c["substance_id"] or None,
                  c["mapping_method"]) for c in comps])
        n_med = db.execute("SELECT COUNT(*) FROM ema_medicine").fetchone()[0]
        n_cmp = db.execute("SELECT COUNT(*) FROM ema_medicine_ingredient").fetchone()[0]
        n_map = db.execute("SELECT COUNT(*) FROM ema_medicine_ingredient "
                           "WHERE substance_id IS NOT NULL").fetchone()[0]
        print(f"=== EMA ingest apply（{args.db}）===")
        print(f"寫入 ema_medicine {n_med:,} 行 / ema_medicine_ingredient {n_cmp:,} 行；"
              f"substance_id 已映射 {n_map:,}（{n_map / max(1, n_cmp) * 100:.1f}%）")
        print(f"region={REGION}、source_id={SOURCE_ID}（原型 staging 表；production 需 region schema）")
        (WORK / "apply_stats.json").write_text(json.dumps(
            {**summary, "db": str(args.db), "applied_at": now,
             "staged_medicines": n_med, "staged_components": n_cmp,
             "staged_components_mapped": n_map}, ensure_ascii=False, indent=1),
            encoding="utf-8")
        return 0
    finally:
        db.close()


def cmd_rollback(args: argparse.Namespace) -> int:
    if not args.db.exists():
        raise SystemExit(f"DB 不存在：{args.db}")
    db = sqlite3.connect(args.db, timeout=120)
    try:
        with db:
            db.executescript(
                "DROP TABLE IF EXISTS ema_medicine_ingredient;"
                "DROP TABLE IF EXISTS ema_medicine;")
        print(f"已刪除 staging 表 ema_medicine / ema_medicine_ingredient（{args.db}）")
        return 0
    finally:
        db.close()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="EMA 集中審批藥品數據（EU 層）入庫")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--db", type=Path, default=DEFAULT_DB)
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="解析+映射+統計（默認）")
    mode.add_argument("--apply", action="store_true", help="寫 staging 原型表（只限副本）")
    mode.add_argument("--rollback", action="store_true", help="刪除 staging 表")
    mode.add_argument("--apply-products", action="store_true",
                      help="寫正式 product/product_component/name（region='EU'；需先遷移 region schema）")
    mode.add_argument("--rollback-products", action="store_true",
                      help="按撤銷清單回滾正式 product 表（--manifest）")
    ap.add_argument("--csv", type=Path, default=WORK / "ema_preview.csv")
    ap.add_argument("--manifest", type=Path, default=PRODUCT_MANIFEST,
                    help="正式表撤銷清單（默認 .tmp/ema-ingest/ema_products_manifest.json）")
    ap.add_argument("--category", choices=["human", "veterinary", "all"], default="human",
                    help="默認 human（藥品站只收人用）；all 含獸藥")
    ap.add_argument("--allow-real", action="store_true",
                    help="允許寫 data/drugs.db（staging 或正式表；默認拒絕）")
    args = ap.parse_args(argv)
    if not args.xlsx.exists():
        raise SystemExit(f"XLSX 不存在：{args.xlsx}")
    if args.rollback:
        return cmd_rollback(args)
    if args.rollback_products:
        return cmd_rollback_products(args)
    if args.apply_products:
        return cmd_apply_products(args)
    if args.apply:
        return cmd_apply(args)
    return cmd_dry_run(args)


if __name__ == "__main__":
    raise SystemExit(main())
